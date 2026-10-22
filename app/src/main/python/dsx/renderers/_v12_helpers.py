"""Helpers for DSX 1.2 features: highlight, comments, cell borders, shapes, list-of."""
from __future__ import annotations

import re
import zipfile
from io import BytesIO
from typing import Any

from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn


def apply_highlight(run: Any, color: str | None) -> None:
    if not color:
        return
    key = color.strip().lstrip("#")
    mapping = {
        "yellow": "yellow", "green": "green", "cyan": "cyan", "magenta": "magenta",
        "blue": "blue", "red": "red", "darkblue": "darkBlue", "darkcyan": "darkCyan",
        "darkgreen": "darkGreen", "darkmagenta": "darkMagenta", "darkred": "darkRed",
        "darkyellow": "darkYellow", "darkgray": "darkGray", "lightgray": "lightGray",
        "black": "black",
    }
    named = mapping.get(key.lower(), "yellow")
    rPr = run._r.get_or_add_rPr()
    for old in rPr.findall(qn("w:highlight")):
        rPr.remove(old)
    hl = OxmlElement("w:highlight")
    hl.set(qn("w:val"), named)
    rPr.append(hl)


def add_comment_markers(par: Any, seg: Any) -> None:
    p = par._p
    crs = OxmlElement("w:commentRangeStart")
    crs.set(qn("w:id"), str(seg.comment_id))
    p.append(crs)
    par.add_run("\u200b")
    cre = OxmlElement("w:commentRangeEnd")
    cre.set(qn("w:id"), str(seg.comment_id))
    p.append(cre)
    ref_run = par.add_run()
    cref = OxmlElement("w:commentReference")
    cref.set(qn("w:id"), str(seg.comment_id))
    ref_run._r.append(cref)


def apply_cell_borders(tc_pr: Any, borders: dict) -> None:
    style_map = {
        "none": "nil", "single": "single", "double": "double",
        "dashed": "dashed", "dotted": "dotted", "thick": "thick",
    }
    existing = tc_pr.find(qn("w:tcBorders"))
    if existing is not None:
        tc_pr.remove(existing)
    tc_borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        info = borders.get(side)
        if not info:
            continue
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), style_map.get(info.get("style", "single"), "single"))
        el.set(qn("w:sz"), str(int(round(float(info.get("width_pt", 0.5)) * 8))))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), str(info.get("color", "000000")).lstrip("#"))
        tc_borders.append(el)
    tc_pr.append(tc_borders)


def build_comments_xml(comments: tuple) -> bytes:
    parts = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">',
    ]

    def esc(s: str) -> str:
        return (
            s.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;")
        )

    for c in comments:
        author = getattr(c, "author", "DSX") or "DSX"
        initials = getattr(c, "initials", "") or ""
        text = getattr(c, "text", "") or ""
        cid = getattr(c, "comment_id", 0)
        parts.append(
            f'<w:comment w:id="{cid}" w:author="{esc(author)}" w:initials="{esc(initials)}" '
            f'w:date="2000-01-01T00:00:00Z">'
            f"<w:p><w:r><w:t>{esc(text)}</w:t></w:r></w:p></w:comment>"
        )
    parts.append("</w:comments>")
    return "\n".join(parts).encode("utf-8")


def inject_native_comments(docx_bytes: bytes, comments: tuple) -> bytes:
    if not comments:
        return docx_bytes
    src = zipfile.ZipFile(BytesIO(docx_bytes), "r")
    out_buf = BytesIO()
    with zipfile.ZipFile(out_buf, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "[Content_Types].xml" and "comments" not in data.decode("utf-8", "replace"):
                s = data.decode("utf-8")
                s = s.replace(
                    "</Types>",
                    '<Override PartName="/word/comments.xml" '
                    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>'
                    "</Types>",
                )
                data = s.encode("utf-8")
            elif info.filename == "word/_rels/document.xml.rels" and "comments.xml" not in data.decode("utf-8", "replace"):
                s = data.decode("utf-8")
                ids = [int(x) for x in re.findall(r'Id="rId(\d+)"', s)]
                next_id = max(ids) + 1 if ids else 1
                s = s.replace(
                    "</Relationships>",
                    f'<Relationship Id="rId{next_id}" '
                    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" '
                    'Target="comments.xml"/></Relationships>',
                )
                data = s.encode("utf-8")
            dst.writestr(info.filename, data)
        dst.writestr("word/comments.xml", build_comments_xml(comments))
    src.close()
    return out_buf.getvalue()


def list_of_block(renderer: Any, doc: Any, block: Any, caption_label: str) -> None:
    title = getattr(block, "title", "") or (
        "Daftar Tabel" if caption_label == "Table" else "Daftar Gambar"
    )
    if title:
        tpar = doc.add_paragraph(title)
        tpar.style = renderer._style_map.get("heading-1", renderer._style_map["normal"])
    par = doc.add_paragraph()
    run = par.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)
    run2 = par.add_run()
    instr = OxmlElement("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = f' TOC \\h \\z \\c "{caption_label}" '
    run2._r.append(instr)
    run3 = par.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run3._r.append(fld_sep)
    par.add_run(f"(Daftar {caption_label.lower()} — perbarui field di Word/LibreOffice)")
    run4 = par.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run4._r.append(fld_end)
    renderer._spacer(doc)


def shape_block(renderer: Any, doc: Any, block: Any) -> None:
    if block.page_break_before:
        doc.add_paragraph().paragraph_format.page_break_before = True
    par = doc.add_paragraph()
    align_map = {"left": 0, "center": 1, "right": 2}
    par.paragraph_format.alignment = align_map.get(block.align, 1)
    cx = int(round(block.width_cm * 360000))
    cy = int(round(block.height_cm * 360000))
    fill_hex = (block.fill or "").lstrip("#") or None
    stroke_hex = (block.stroke or "000000").lstrip("#")
    stroke_w = int(round(block.stroke_width_pt * 12700))
    geom_map = {"rect": "rect", "ellipse": "ellipse", "roundRect": "roundRect", "line": "line"}
    geom = geom_map.get(block.shape_type, "rect")
    if fill_hex:
        fill_xml = f'<a:solidFill><a:srgbClr val="{fill_hex}"/></a:solidFill>'
    else:
        fill_xml = "<a:noFill/>"
    if block.stroke_width_pt > 0:
        stroke_xml = (
            f'<a:ln w="{stroke_w}">'
            f'<a:solidFill><a:srgbClr val="{stroke_hex}"/></a:solidFill></a:ln>'
        )
    else:
        stroke_xml = "<a:ln><a:noFill/></a:ln>"
    adj = ""
    if block.shape_type == "roundRect":
        pct = min(50000, int(round((block.corner_cm / max(block.width_cm, 0.01)) * 100000)))
        adj = f'<a:avLst><a:gd name="adj" fmla="val {pct}"/></a:avLst>'
    _wpns = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
    _used = {int(e.get("id")) for e in doc.element.body.iter(_wpns) if str(e.get("id", "")).isdigit()}
    docpr_id = max(_used, default=0) + 1
    shape_xml = (
        '<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wp:inline distT="0" distB="0" distL="0" distR="0">'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        f'<wp:docPr id="{docpr_id}" name="Shape {docpr_id}"/>'
        "<a:graphic><a:graphicData "
        'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        "<wps:wsp><wps:cNvCnPr/><wps:spPr>"
        f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="{geom}">{adj}</a:prstGeom>'
        f"{fill_xml}{stroke_xml}"
        "</wps:spPr><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic>"
        "</wp:inline></w:drawing>"
    )
    try:
        drawing = parse_xml(shape_xml)
        run = par.add_run()
        run._r.append(drawing)
    except Exception:
        par.add_run(f"[SHAPE {block.shape_type} {block.width_cm}x{block.height_cm}cm]")
